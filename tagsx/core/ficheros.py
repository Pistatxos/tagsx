"""Leer y escribir el inventario en CSV y XLSX, con los arreglos de Excel en español."""
import csv
import io
import os
import tempfile


def read_csv_rows(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return _leer_csv(f)


def _leer_csv(f):
    # encoding="utf-8-sig" porque Excel suele anadir un BOM al guardar CSV.
    # delimiters=",;" porque el Excel en español exporta CSV con ";" (usa
    # "," como separador decimal) - sin esto, una columna con comas dentro
    # se puede partir en columnas de mas y el script no lo detecta bien.
    sample = f.read(4096)
    f.seek(0)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(f, dialect=dialect)
    fieldnames = reader.fieldnames or []
    rows = [dict(row) for row in reader]
    return fieldnames, rows


def read_xlsx_rows(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            return [], []
        fieldnames = [str(h).strip() if h is not None else "" for h in header]

        rows = []
        for values in rows_iter:
            if all(v is None for v in values):
                continue
            row = {}
            for key, value in zip(fieldnames, values):
                if value is None:
                    row[key] = ""
                elif isinstance(value, bool):
                    # Excel guarda TRUE/FALSE (VERDADERO/FALSO) como celda
                    # booleana, no como texto. openpyxl la devuelve como
                    # bool nativo de Python: sin esto str(value) daria
                    # "True"/"False" con mayuscula, distinto de los valores
                    # 'true'/'false' en minuscula que AWS usa de verdad
                    # (p.ej. tags PatchEnabled de SSM o dlm:managed de DLM).
                    row[key] = "true" if value else "false"
                elif isinstance(value, float) and value.is_integer():
                    row[key] = str(int(value))
                else:
                    row[key] = str(value)
            rows.append(row)
        return fieldnames, rows
    finally:
        wb.close()


def read_input_rows(path):
    if path.lower().endswith(".xlsx"):
        return read_xlsx_rows(path)
    return read_csv_rows(path)


def leer_subido(nombre, contenido):
    """Un fichero subido desde la web (bytes)."""
    if nombre.lower().endswith(".xlsx"):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as t:
            t.write(contenido)
        try:
            return read_xlsx_rows(t.name)
        finally:
            os.unlink(t.name)
    texto = contenido.decode("utf-8-sig", errors="replace")
    return _leer_csv(io.StringIO(texto, newline=""))


def escribir_csv(destino, fieldnames, filas):
    """`destino`: una ruta o un objeto de texto abierto."""
    if isinstance(destino, str):
        with open(destino, "w", newline="", encoding="utf-8") as f:
            return escribir_csv(f, fieldnames, filas)
    writer = csv.DictWriter(destino, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(filas)


def escribir_xlsx(destino, fieldnames, filas, columnas_info=()):
    """Todo como texto (formato @), para que Excel no convierta true/false, números largos
    o fechas. Cabecera fija y con filtro; las columnas informativas, en gris."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "inventario"
    ws.append(fieldnames)
    for fila in filas:
        ws.append([fila.get(c, "") for c in fieldnames])
    gris = PatternFill("solid", fgColor="E9EFF5")
    for i, nombre in enumerate(fieldnames, start=1):
        letra = get_column_letter(i)
        ws.cell(row=1, column=i).font = Font(bold=True)
        if nombre in columnas_info:
            ws.cell(row=1, column=i).fill = gris
        largo = max([len(nombre)] + [len(str(f.get(nombre, ""))) for f in filas[:500]])
        ws.column_dimensions[letra].width = min(max(10, largo + 2), 60)
        for celda in ws[letra][1:]:
            celda.number_format = "@"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(destino)
